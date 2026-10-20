import io
import unittest

from openpyxl import load_workbook

from excel_handler import write_excel
from solver import HOLIDAY, _normalize_cell, generate_shift

STAFF_COUNT = 19
DAY_COUNT = 30
SETTINGS = {
    "year": 2026,
    "month": 11,
    "day_leader_count": 3,
    "night_leader_count": 8,
    "night_eligible_count": 17,
    "max_night_shifts": 5,
}


def staff_ids():
    return [str(i + 1) for i in range(STAFF_COUNT)]


def blank_schedule():
    return [[""] * DAY_COUNT for _ in range(STAFF_COUNT)]


def generate(schedule):
    return generate_shift(
        staff_ids(), 2026, 11, schedule, SETTINGS
    )


class HolidayLabelTests(unittest.TestCase):
    def test_all_holiday_spellings_normalize_to_holiday(self):
        self.assertEqual(HOLIDAY, "公休")
        for text in ("公", "休", "公休"):
            self.assertEqual(_normalize_cell(text), "公休")

    def test_generated_schedule_uses_holiday_label_only(self):
        result, _, _ = generate(blank_schedule())
        cells = [cell for row in result for cell in row]
        self.assertNotIn("公", cells)
        self.assertIn("公休", cells)

    def test_legacy_fixed_holiday_becomes_holiday_label(self):
        schedule = blank_schedule()
        schedule[0][4] = "公"
        schedule[1][5] = "休"
        result, _, _ = generate(schedule)
        self.assertEqual(result[0][4], "公休")
        self.assertEqual(result[1][5], "公休")
        self.assertNotIn("公", [cell for row in result for cell in row])


class WriteExcelHolidayTests(unittest.TestCase):
    def test_holiday_is_green_and_requested_off_is_pink(self):
        schedule = [["公休", "公休", "公休", "公休", "公休", "日"]]
        original = [["公", "休", "公休", "希", "", "日"]]
        data = write_excel(["1"], list(range(1, 7)), schedule, None, original)
        ws = load_workbook(io.BytesIO(data)).active

        def fill(col):
            return ws.cell(2, col).fill.start_color.rgb[-6:]

        self.assertEqual(ws.cell(2, 6).value, "公休")
        for col in (2, 3, 4, 5):
            self.assertEqual(fill(col), "fce7f3")
        self.assertEqual(fill(6), "c8e6c9")
        self.assertNotEqual(fill(7), "fce7f3")


if __name__ == "__main__":
    unittest.main()

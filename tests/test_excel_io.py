import asyncio
import io
import unittest
from unittest import mock

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from starlette.datastructures import UploadFile

from excel_handler import read_excel, write_excel
from main import DownloadRequest, download_excel, upload_excel

XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def build_xlsx(rows):
    """行のリストから xlsx のバイト列を合成する（職員番号は連番、実名は使わない）"""
    wb = Workbook()
    ws = wb.active
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row, start=1):
            if val is not None:
                ws.cell(r, c, val)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def upload(contents, filename="test.xlsx"):
    """upload_excel を直接呼び、(status_code, detail または戻り値) を返す"""
    file = UploadFile(file=io.BytesIO(contents), filename=filename)
    try:
        return 200, asyncio.run(upload_excel(file))
    except HTTPException as e:
        return e.status_code, e.detail


def open_written(**kwargs):
    return load_workbook(io.BytesIO(write_excel(**kwargs))).active


def fill_of(cell):
    return cell.fill.start_color.rgb[-6:].lower()


class ReadExcelTests(unittest.TestCase):
    def test_basic_layout(self):
        data = build_xlsx([
            ["職員番号", 1, 2, 3],
            ["001", "日", "", "公休"],
            ["002", "", "夜", "明"],
        ])
        self.assertEqual(read_excel(data), {
            "staff_ids": ["001", "002"],
            "dates": [1, 2, 3],
            "schedule": [["日", "", "公休"], ["", "夜", "明"]],
        })

    def test_date_start_column_offset(self):
        # 日付 1 が C 列から始まる（B 列は別の見出し）。シフトも C 列から読む
        data = build_xlsx([
            ["職員番号", "区分", 1, 2],
            ["001", "x", "日", "夜"],
        ])
        result = read_excel(data)
        self.assertEqual(result["dates"], [1, 2])
        self.assertEqual(result["schedule"], [["日", "夜"]])

    def test_date_header_as_text_and_float(self):
        data = build_xlsx([["職員番号", "1", 2.0, 3], ["001", "日", "日", "日"]])
        self.assertEqual(read_excel(data)["dates"], [1, 2, 3])

    def test_date_header_stops_at_gap_or_non_sequence(self):
        # 連番が途切れたらそこまでを日付とする（曜日行などを巻き込まない）
        data = build_xlsx([
            ["職員番号", 1, 2, 4, 5],
            ["001", "日", "日", "日", "日"],
        ])
        result = read_excel(data)
        self.assertEqual(result["dates"], [1, 2])
        self.assertEqual(result["schedule"], [["日", "日"]])

    def test_date_header_stops_after_31(self):
        data = build_xlsx([
            ["職員番号"] + list(range(1, 33)),
            ["001"] + ["日"] * 32,
        ])
        self.assertEqual(read_excel(data)["dates"], list(range(1, 32)))

    def test_empty_cells_become_empty_string(self):
        data = build_xlsx([["職員番号", 1, 2, 3], ["001", None, "日", None]])
        self.assertEqual(read_excel(data)["schedule"], [["", "日", ""]])

    def test_whitespace_is_stripped_including_fullwidth(self):
        data = build_xlsx([
            ["職員番号", 1, 2],
            ["  001 ", " 日 ", "公休　"],
        ])
        result = read_excel(data)
        self.assertEqual(result["staff_ids"], ["001"])
        self.assertEqual(result["schedule"], [["日", "公休"]])

    def test_numeric_staff_id_becomes_string(self):
        data = build_xlsx([["職員番号", 1], [1, "日"], [2, "夜"]])
        self.assertEqual(read_excel(data)["staff_ids"], ["1", "2"])

    def test_blank_rows_before_first_staff_are_skipped(self):
        data = build_xlsx([
            ["職員番号", 1],
            [None, None],
            ["001", "日"],
        ])
        self.assertEqual(read_excel(data)["staff_ids"], ["001"])

    def test_blank_row_after_staff_ends_the_list(self):
        # 空行以降の凡例・集計は職員として読まない
        data = build_xlsx([
            ["職員番号", 1],
            ["001", "日"],
            ["002", "夜"],
            [None, None],
            ["凡例", "日"],
        ])
        result = read_excel(data)
        self.assertEqual(result["staff_ids"], ["001", "002"])
        self.assertEqual(len(result["schedule"]), 2)

    def test_trailing_blank_rows_are_ignored(self):
        data = build_xlsx([["職員番号", 1], ["001", "日"], [None, None], [None, None]])
        self.assertEqual(read_excel(data)["staff_ids"], ["001"])

    def test_missing_date_header_raises(self):
        data = build_xlsx([["職員番号", "日付"], ["001", "日"]])
        with self.assertRaises(ValueError) as ctx:
            read_excel(data)
        self.assertIn("日付 '1'", str(ctx.exception))

    def test_header_without_staff_rows_returns_empty(self):
        result = read_excel(build_xlsx([["職員番号", 1, 2]]))
        self.assertEqual(result["staff_ids"], [])
        self.assertEqual(result["schedule"], [])


class UploadErrorTests(unittest.TestCase):
    # パスワード付き・.xls・壊れたバイト列は tests/test_unsupported_excel.py で検証済み

    def test_missing_date_header_returns_400_with_prefix(self):
        status, detail = upload(build_xlsx([["職員番号", "日付"], ["001", "日"]]))
        self.assertEqual(status, 400)
        self.assertEqual(
            detail,
            "ファイルの読み込みに失敗しました: "
            "1行目で日付 '1' が入っている列を見つけられませんでした。",
        )

    def test_no_staff_returns_400(self):
        status, detail = upload(build_xlsx([["職員番号", 1, 2]]))
        self.assertEqual(status, 400)
        self.assertEqual(
            detail, "職員データが見つかりません。フォーマットを確認してください。"
        )

    def test_non_xlsx_extension_returns_400(self):
        status, detail = upload(build_xlsx([["職員番号", 1], ["001", "日"]]), "data.csv")
        self.assertEqual(status, 400)
        self.assertEqual(detail, "xlsx形式のファイルをアップロードしてください")

    def test_valid_file_returns_parsed_data(self):
        status, body = upload(build_xlsx([["職員番号", 1, 2], ["001", "日", "夜"]]))
        self.assertEqual(status, 200)
        self.assertEqual(body["dates"], [1, 2])
        self.assertEqual(body["schedule"], [["日", "夜"]])


class WriteExcelTests(unittest.TestCase):
    def test_cell_values_and_header(self):
        ws = open_written(
            staff_ids=["001", "002"],
            dates=[1, 2],
            schedule=[["日", "公休"], ["夜", "明"]],
        )
        self.assertEqual(ws.title, "勤務表")
        self.assertEqual(ws["A1"].value, "職員番号")
        self.assertEqual([ws["B1"].value, ws["C1"].value], [1, 2])
        self.assertEqual(ws["A2"].value, "001")
        self.assertEqual([ws["B2"].value, ws["C2"].value], ["日", "公休"])
        self.assertEqual([ws["B3"].value, ws["C3"].value], ["夜", "明"])
        self.assertEqual(ws.freeze_panes, "B2")

    def test_column_widths(self):
        ws = open_written(staff_ids=["001"], dates=[1, 2], schedule=[["日", "夜"]])
        self.assertEqual(ws.column_dimensions["A"].width, 12)
        self.assertEqual(ws.column_dimensions["B"].width, 4.5)
        self.assertEqual(ws.column_dimensions["C"].width, 4.5)

    def test_shift_colors(self):
        shifts = ["日", "夜", "明", "公", "公休", "希", "委", "休", "有", "研"]
        ws = open_written(
            staff_ids=["001"], dates=list(range(1, 11)), schedule=[shifts]
        )
        expected = {
            "日": "ffffff", "夜": "1e3a5f", "明": "bbdefb", "公": "c8e6c9",
            "公休": "c8e6c9", "希": "fce7f3", "委": "ffe0b2", "休": "e0e0e0",
            "有": "e0e0e0", "研": "e1bee7",
        }
        for i, shift in enumerate(shifts):
            self.assertEqual(fill_of(ws.cell(2, i + 2)), expected[shift], shift)

    def test_night_shift_is_bold_white_text(self):
        ws = open_written(staff_ids=["001"], dates=[1, 2], schedule=[["夜", "日"]])
        self.assertTrue(ws["B2"].font.bold)
        self.assertEqual(ws["B2"].font.color.rgb[-6:].upper(), "FFFFFF")
        self.assertFalse(ws["C2"].font.bold)

    def test_unknown_or_empty_shift_has_no_fill(self):
        ws = open_written(staff_ids=["001"], dates=[1, 2], schedule=[["", "謎"]])
        self.assertIsNone(ws["B2"].fill.fill_type)
        self.assertIsNone(ws["C2"].fill.fill_type)

    def test_original_rest_cells_get_pink(self):
        # 元データが休み系のセルは、生成後に「公休」でも希望休のピンクにする
        original = [["公", "公休", "希", "休", "希望休", "希望", "日", ""]]
        ws = open_written(
            staff_ids=["001"], dates=list(range(1, 9)),
            schedule=[["公休"] * 8], original_schedule=original,
        )
        for col in range(2, 8):
            self.assertEqual(fill_of(ws.cell(2, col)), "fce7f3", col)
        # 休み系でない元データ（日・空）は公休の薄緑のまま
        self.assertEqual(fill_of(ws.cell(2, 8)), "c8e6c9")
        self.assertEqual(fill_of(ws.cell(2, 9)), "c8e6c9")

    def test_original_schedule_shorter_than_staff_is_safe(self):
        ws = open_written(
            staff_ids=["001", "002"], dates=[1], schedule=[["公休"], ["公休"]],
            original_schedule=[["希"]],
        )
        self.assertEqual(fill_of(ws["B2"]), "fce7f3")
        self.assertEqual(fill_of(ws["B3"]), "c8e6c9")

    def test_day_leader_fill_only_and_text_unchanged(self):
        ws = open_written(
            staff_ids=["001", "002"], dates=[1, 2],
            schedule=[["日", "日"], ["日", "日"]],
            day_leaders=["001", "002"],
        )
        self.assertEqual(fill_of(ws["B2"]), "fff2cc")
        self.assertEqual(fill_of(ws["C3"]), "fff2cc")
        self.assertEqual(fill_of(ws["C2"]), "ffffff")
        self.assertEqual(ws["B2"].value, "日")

    def test_day_leader_overrides_pink(self):
        ws = open_written(
            staff_ids=["001"], dates=[1], schedule=[["日"]],
            original_schedule=[["希"]], day_leaders=["001"],
        )
        self.assertEqual(fill_of(ws["B2"]), "fff2cc")

    def test_day_leaders_none_entry_and_short_list(self):
        ws = open_written(
            staff_ids=["001"], dates=[1, 2], schedule=[["日", "日"]],
            day_leaders=[None],
        )
        self.assertEqual(fill_of(ws["B2"]), "ffffff")
        self.assertEqual(fill_of(ws["C2"]), "ffffff")


class RoundTripTests(unittest.TestCase):
    def test_write_then_read_matches(self):
        # 書き出しは A1 が「職員番号」、日付は B1 から始まり、読み込み形式と同じ。
        # 差は装飾（色・罫線・列幅・フリーズ）だけで、読み込み時には無視される。
        staff_ids = ["001", "002", "003"]
        dates = list(range(1, 8))
        schedule = [
            ["日", "日", "公休", "夜", "明", "公休", "公休"],
            ["", "有", "研", "委", "休", "希", "日"],
            ["夜", "明", "公休", "日", "日", "日", "公休"],
        ]
        data = write_excel(staff_ids, dates, schedule, day_leaders=[None] * 7)
        result = read_excel(data)
        self.assertEqual(result["staff_ids"], staff_ids)
        self.assertEqual(result["dates"], dates)
        self.assertEqual(result["schedule"], schedule)

    def test_roundtrip_through_upload(self):
        status, body = upload(write_excel(["001"], [1, 2], [["日", "夜"]]))
        self.assertEqual(status, 200)
        self.assertEqual(body["schedule"], [["日", "夜"]])


class DownloadExcelTests(unittest.TestCase):
    def request(self, **overrides):
        params = dict(
            staff_ids=["001", "002"], dates=[1, 2],
            schedule=[["日", "公休"], ["夜", "明"]],
        )
        params.update(overrides)
        return DownloadRequest(**params)

    def test_success_returns_openable_xlsx_with_headers(self):
        response = asyncio.run(download_excel(self.request(day_leaders=["001", None])))
        self.assertEqual(response.media_type, XLSX_TYPE)
        self.assertEqual(
            response.headers["content-disposition"],
            'attachment; filename="shift_schedule.xlsx"',
        )
        ws = load_workbook(io.BytesIO(response.body)).active
        self.assertEqual(ws["B3"].value, "夜")
        self.assertEqual(fill_of(ws["B2"]), "fff2cc")

    def test_generation_failure_returns_500(self):
        # schedule の行数が職員数より少ないと write_excel が IndexError になる
        with self.assertRaises(HTTPException) as ctx:
            asyncio.run(download_excel(self.request(schedule=[["日", "夜"]])))
        self.assertEqual(ctx.exception.status_code, 500)
        self.assertTrue(
            ctx.exception.detail.startswith("エクセル生成中にエラーが発生しました: ")
        )

    def test_write_excel_exception_message_is_included(self):
        with mock.patch("main.write_excel", side_effect=RuntimeError("boom")):
            with self.assertRaises(HTTPException) as ctx:
                asyncio.run(download_excel(self.request()))
        self.assertEqual(ctx.exception.status_code, 500)
        self.assertEqual(
            ctx.exception.detail, "エクセル生成中にエラーが発生しました: boom"
        )


if __name__ == "__main__":
    unittest.main()

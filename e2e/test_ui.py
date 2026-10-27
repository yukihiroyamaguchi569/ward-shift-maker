"""画面（static/app.js）の E2E テスト。

実サーバー（uvicorn）を起動し、ヘッドレス Chromium で操作する。
入力の xlsx はテスト内で合成する（実データは使わない）。
"""
import os
import re
import socket
import subprocess
import sys
import tempfile
import time
import unittest
import urllib.request
from io import BytesIO
from urllib.parse import parse_qs, urlparse

from openpyxl import Workbook
from playwright.sync_api import expect, sync_playwright

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
DAY_COUNT = 30  # 2026年11月
MONTH = "2026-11"
SERVER_START_TIMEOUT_SECONDS = 30
DIALOG_WAIT_SECONDS = 30
UI_TIMEOUT_MS = 60000  # 生成は制限時間 10 秒＋余裕

PRECHECK_MESSAGE = "このままでは勤務表を作成できません。"
PASSWORD_WARNING = "Excel にパスワードがかかっていると読み込めません。"
# フィードバック用フォームの仮の設定（本物の ID が入る前でもリンクの中身を確かめるため）
FAKE_FEEDBACK_FORM = '''const FEEDBACK_FORM = {
    url: "https://docs.google.com/forms/d/e/FAKE/viewform",
    entries: {
        operation: "entry.1",
        version: "entry.2",
        stage: "entry.3",
        error: "entry.4",
        userAgent: "entry.5",
        datetime: "entry.6",
    },
};'''
FEEDBACK_FIELDS = {
    "entry.1": "operation",
    "entry.2": "version",
    "entry.3": "stage",
    "entry.4": "error",
    "entry.5": "userAgent",
    "entry.6": "datetime",
}
PASSWORD_BYTES = (
    bytes.fromhex("D0CF11E0A1B11AE1")
    + b"\x00" * 16
    + "EncryptionInfo".encode("utf-16-le")
)


def free_port():
    with socket.socket() as s:
        s.bind(("127.0.0.1", 0))
        return s.getsockname()[1]


def make_xlsx(staff_count, fixed_cells=None):
    """1行目に日付 1..30、A列に職員番号の勤務表を作る。fixed_cells は {(職員index, 日index): 値}"""
    wb = Workbook()
    ws = wb.active
    ws.cell(1, 1, "職員番号")
    for d in range(DAY_COUNT):
        ws.cell(1, 2 + d, d + 1)
    for s in range(staff_count):
        ws.cell(2 + s, 1, str(s + 1))
    for (s, d), value in (fixed_cells or {}).items():
        ws.cell(2 + s, 2 + d, value)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def overbooked_day_29():
    # 17人中、29日に 委 が1人、公 が11人 → 配置可能が足りない（tests/test_precheck.py と同じ入力）
    cells = {(0, 28): "委"}
    for s in range(1, 12):
        cells[(s, 28)] = "公"
    return make_xlsx(17, cells)


def day_leader_blocked():
    # 10日: 日勤リーダー候補（上から3人）が 公・委・休 で、誰も日勤に入れない（tests/test_diagnosis.py と同じ入力）
    return make_xlsx(19, {(0, 9): "公", (1, 9): "委", (2, 9): "休"})


class UiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        port = free_port()
        cls.base_url = f"http://127.0.0.1:{port}"
        env = {**os.environ, "SOLVER_TIME_LIMIT_SECONDS": "10"}
        cls.server = subprocess.Popen(
            [sys.executable, "-m", "uvicorn", "main:app", "--port", str(port)],
            cwd=ROOT,
            env=env,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
        )
        cls.playwright = None
        cls.browser = None
        # setUpClass が途中で失敗しても tearDownClass は呼ばれないため、ここで後始末を登録する
        cls.addClassCleanup(cls.stop_all)
        cls.wait_for_server()
        cls.playwright = sync_playwright().start()
        cls.browser = cls.playwright.chromium.launch()

    @classmethod
    def wait_for_server(cls):
        deadline = time.monotonic() + SERVER_START_TIMEOUT_SECONDS
        while time.monotonic() < deadline:
            if cls.server.poll() is not None:
                raise RuntimeError("uvicorn が起動直後に終了しました")
            try:
                urllib.request.urlopen(cls.base_url + "/", timeout=2).close()
                return
            except OSError:
                time.sleep(0.2)
        raise RuntimeError("uvicorn が時間内に応答しませんでした")

    @classmethod
    def stop_all(cls):
        # ブラウザの終了で例外が出ても、サーバーは必ず止める
        try:
            if cls.browser:
                cls.browser.close()
            if cls.playwright:
                cls.playwright.stop()
        finally:
            cls.server.terminate()
            try:
                cls.server.wait(timeout=10)
            except subprocess.TimeoutExpired:
                cls.server.kill()
                cls.server.wait()

    def setUp(self):
        # テストごとに新しいコンテキスト（localStorage を持ち越さない）
        self.context = self.browser.new_context(accept_downloads=True)
        self.addCleanup(self.context.close)
        # 外部 CDN（Tailwind・フォント）には出ない。画面の動作に必要なのは自サーバーの資源だけ
        self.context.route(
            "**/*",
            lambda route: route.continue_()
            if route.request.url.startswith(self.base_url)
            else route.abort(),
        )
        self.context.set_default_timeout(UI_TIMEOUT_MS)
        self.page = self.context.new_page()
        self.dialogs = []
        self.prompt_answer = ""
        self.page.on("dialog", self.on_dialog)
        self.page.goto(self.base_url + "/", wait_until="domcontentloaded")
        self.page.fill("#targetMonth", MONTH)
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)

    def on_dialog(self, dialog):
        self.dialogs.append(dialog.message)
        if dialog.type == "prompt":
            dialog.accept(self.prompt_answer)
        else:
            dialog.accept()

    def upload(self, contents, filename="勤務表.xlsx"):
        self.page.set_input_files(
            "#fileInput",
            files=[{"name": filename, "mimeType": "application/octet-stream", "buffer": contents}],
        )

    def warnings_text(self):
        return self.page.locator("#warningsList").inner_text()

    def wait_for_dialog(self):
        deadline = time.monotonic() + DIALOG_WAIT_SECONDS
        while not self.dialogs and time.monotonic() < deadline:
            self.page.wait_for_timeout(100)
        self.assertEqual(len(self.dialogs), 1, self.dialogs)

    def use_fake_feedback_form(self):
        """app.js の FEEDBACK_FORM を仮の設定に差し替えて読み込み直す"""
        with open(os.path.join(ROOT, "static", "app.js"), encoding="utf-8") as f:
            app_js = f.read()
        start = app_js.index("const FEEDBACK_FORM = {")
        end = app_js.index("};", start) + 2
        patched = app_js[:start] + FAKE_FEEDBACK_FORM + app_js[end:]
        self.page.route(
            "**/static/app.js*",
            lambda route: route.fulfill(body=patched, content_type="application/javascript"),
        )
        self.page.reload(wait_until="domcontentloaded")
        self.page.fill("#targetMonth", MONTH)

    def feedback_values(self):
        """リンクを押す直前と同じように URL を作り直させ、事前入力の値を返す"""
        self.page.dispatch_event("#feedbackLink", "pointerdown")
        query = parse_qs(urlparse(self.page.get_attribute("#feedbackLink", "href")).query)
        self.assertEqual(query.pop("usp"), ["pp_url"])
        return {FEEDBACK_FIELDS[k]: v[0] for k, v in query.items()}

    def test_feedback_link_hidden_until_form_is_configured(self):
        expect(self.page.locator("#feedbackLink")).to_be_hidden()

    def test_feedback_link_prefills_state(self):
        self.use_fake_feedback_form()
        link = self.page.locator("#feedbackLink")
        expect(link).to_be_visible()
        self.assertEqual(link.get_attribute("target"), "_blank")
        values = self.feedback_values()
        self.assertEqual(values["stage"], "ファイル未アップロード")
        self.assertRegex(values["version"], r"^\d+\.\d+\.\d+$")
        self.assertRegex(values["datetime"], r"^\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2}$")
        self.assertIn("userAgent", values)
        self.assertNotIn("error", values)
        # setUp / use_fake_feedback_form が対象年月を変えるので、最後の操作は設定の変更
        self.assertEqual(values["operation"], "設定・プリセット")

        # 失敗したアップロードのエラーが入り、次に成功したら消える
        self.upload(b"x", "メモ.txt")
        self.wait_for_dialog()
        values = self.feedback_values()
        self.assertEqual(values["operation"], "ファイルのアップロード")
        self.assertEqual(values["error"], "xlsx形式のファイルを選択してください")

        self.upload(make_xlsx(19))
        self.page.locator("#staffCount", has_text="19名 × 30日").wait_for()
        values = self.feedback_values()
        self.assertEqual(values["operation"], "ファイルのアップロード")
        self.assertEqual(values["stage"], "ファイル読み込み済み（シフト未作成）")
        self.assertNotIn("error", values)

        self.page.click("#generateBtn")
        self.page.locator("#generateBtn", has_text="再作成").wait_for()
        values = self.feedback_values()
        self.assertEqual(values["operation"], "シフト作成")
        self.assertEqual(values["stage"], "シフト作成済み")
        # エラー欄は、いま画面に出ている警告と同じ（無ければ空）
        panel_hidden = "panel-hidden" in self.page.get_attribute("#warningsPanel", "class")
        on_screen = "" if panel_hidden else " ".join(self.page.locator("#warningsList li").all_inner_texts())
        if on_screen:
            self.assertTrue(on_screen.startswith(values["error"].rstrip("…")), values["error"])
        else:
            self.assertNotIn("error", values)

    def test_feedback_link_records_precheck_conflict_after_settings_change(self):
        self.use_fake_feedback_form()
        self.upload(overbooked_day_29())
        self.page.locator("#warningsList li", has_text=PRECHECK_MESSAGE).wait_for()
        self.assertEqual(self.feedback_values()["operation"], "ファイルのアップロード")

        self.page.fill("#maxNightShifts", "4")
        # 設定を変えると操作は「設定・プリセット」になり、作成前チェックの結果が入り直す
        self.page.wait_for_function(
            "document.getElementById('presetSelect').value === '__custom__'"
        )
        self.page.locator("#warningsList li", has_text=PRECHECK_MESSAGE).wait_for()
        expect(self.page.locator("#feedbackLink")).to_be_visible()
        deadline = time.monotonic() + DIALOG_WAIT_SECONDS
        while time.monotonic() < deadline:
            values = self.feedback_values()
            if "error" in values:
                break
            self.page.wait_for_timeout(100)
        self.assertEqual(values["operation"], "設定・プリセット")
        self.assertIn(PRECHECK_MESSAGE, values["error"])
        self.assertIn("29日", values["error"])

    def test_feedback_link_keeps_failed_operation_when_settings_change_meanwhile(self):
        self.use_fake_feedback_form()
        self.upload(make_xlsx(19))
        self.page.locator("#staffCount", has_text="19名 × 30日").wait_for()
        pending = []
        self.page.route("**/api/generate", lambda route: pending.append(route))
        self.page.click("#generateBtn")
        deadline = time.monotonic() + DIALOG_WAIT_SECONDS
        while not pending and time.monotonic() < deadline:
            self.page.wait_for_timeout(100)
        self.assertEqual(len(pending), 1)
        # 作成の応答を待つ間に設定を変える
        self.page.fill("#maxNightShifts", "4")
        self.assertEqual(self.feedback_values()["operation"], "設定・プリセット")
        pending[0].fulfill(status=500, content_type="application/json", body='{"detail": "サーバーエラーです"}')
        self.wait_for_dialog()
        values = self.feedback_values()
        self.assertEqual(values["operation"], "シフト作成")
        self.assertEqual(values["error"], "サーバーエラーです")

    def test_feedback_link_records_generation_warnings(self):
        self.use_fake_feedback_form()
        self.page.fill("#dayLeaderCount", "3")
        self.upload(day_leader_blocked())
        self.page.click("#generateBtn")
        self.page.locator(
            "#warningsList li", has_text="条件を満たす勤務表が見つかりませんでした。"
        ).wait_for()
        values = self.feedback_values()
        self.assertEqual(values["operation"], "シフト作成")
        self.assertIn("条件を満たす勤務表が見つかりませんでした。", values["error"])
        self.assertLessEqual(len(values["error"]), 300)

        # 設定を変えても警告は画面に残るので、エラー欄にも残る
        self.page.fill("#maxNightShifts", "4")
        values = self.feedback_values()
        self.assertEqual(values["operation"], "設定・プリセット")
        self.assertIn("条件を満たす勤務表が見つかりませんでした。", values["error"])

    def test_upload_generate_and_download(self):
        self.upload(make_xlsx(19))
        self.page.locator("#staffCount", has_text="19名 × 30日").wait_for()
        # 職員 19 行＋日勤計・夜勤計の 2 行
        self.assertEqual(self.page.locator("#scheduleTable tbody tr").count(), 19 + 2)
        self.assertNotIn("公休", self.page.locator("#scheduleTable tbody").inner_text())

        self.page.click("#generateBtn")
        self.page.locator("#generateBtn", has_text="再作成").wait_for()
        self.assertIn("公休", self.page.locator("#scheduleTable tbody").inner_text())
        self.assertTrue(self.page.locator("#downloadBtn").is_enabled())

        with self.page.expect_download() as info:
            self.page.click("#downloadBtn")
        path = os.path.join(self.tmp.name, "out.xlsx")
        info.value.save_as(path)
        with open(path, "rb") as f:
            self.assertEqual(f.read(2), b"PK")  # xlsx は zip

    def test_precheck_blocks_overbooked_day(self):
        self.upload(overbooked_day_29())
        self.page.locator("#warningsList li", has_text=PRECHECK_MESSAGE).wait_for()
        self.assertIn("29日", self.warnings_text())

        # そのまま作成を押しても勤務表は作られない（競合の理由が出て、空欄は補完されない）
        self.page.click("#generateBtn")
        self.page.locator("#generateBtn", has_text="再作成").wait_for()
        self.assertIn("29日", self.warnings_text())
        self.assertNotIn(PRECHECK_MESSAGE, self.warnings_text())
        # 入力が全部空欄の最終職員の行は、日勤・夜勤・公休が入らず空のまま
        last_row_cells = self.page.locator("#scheduleTable tbody tr:nth-child(17) td.shift-cell")
        self.assertEqual(last_row_cells.count(), 30)
        self.assertEqual(set(last_row_cells.all_inner_texts()), {""})
        # 現状の挙動の記録: 競合で未完成の表でも、列数が合えばダウンロードは有効になる
        self.assertTrue(self.page.locator("#downloadBtn").is_enabled())

    def test_diagnosis_shows_cause_when_unsolvable(self):
        self.page.fill("#dayLeaderCount", "3")
        self.upload(day_leader_blocked())
        self.page.click("#generateBtn")
        self.page.locator(
            "#warningsList li", has_text="10日: 日勤リーダーを置けません。"
        ).wait_for()
        self.assertIn("条件を満たす勤務表が見つかりませんでした。", self.warnings_text())
        self.assertTrue(self.page.locator("#generateBtn").is_enabled())

    def test_password_protected_file_shows_message(self):
        self.upload(PASSWORD_BYTES, "パスワード付き.xlsx")
        # 読み込み失敗は alert で知らせる
        self.wait_for_dialog()
        self.assertIn(PASSWORD_WARNING, self.dialogs[0])
        self.assertTrue(self.page.locator("#generateBtn").is_disabled())

    def assert_generate_failure_recovers(self, handler):
        self.upload(make_xlsx(19))
        self.page.locator("#staffCount", has_text="19名 × 30日").wait_for()
        self.page.route("**/api/generate", handler)
        self.page.click("#generateBtn")
        self.wait_for_dialog()
        # 固まらず、もう一度押せる
        # alert は finally より先に出ることがあるので、復帰は expect で待つ
        expect(self.page.locator("#loading")).to_have_class(re.compile(r"panel-hidden"))
        expect(self.page.locator("#generateBtn")).to_be_enabled()
        self.assertEqual(self.page.locator("#generateBtn").inner_text(), "シフト作成")
        self.assertTrue(self.page.locator("#downloadBtn").is_disabled())

    def test_generate_http_500_recovers(self):
        self.assert_generate_failure_recovers(
            lambda route: route.fulfill(
                status=500,
                content_type="application/json",
                body='{"detail": "サーバーエラーです"}',
            )
        )
        self.assertIn("サーバーエラーです", self.dialogs[0])

    def test_generate_network_failure_recovers(self):
        self.assert_generate_failure_recovers(lambda route: route.abort())

    def test_sunday_day_required_preset_roundtrip(self):
        self.prompt_answer = "日曜4人"
        self.page.fill("#sundayDayRequired", "4")
        self.page.click("#presetSaveBtn")
        self.assertEqual(self.page.locator("#presetSelect").input_value(), "日曜4人")

        self.page.select_option("#presetSelect", "デフォルト")
        self.assertEqual(self.page.locator("#sundayDayRequired").input_value(), "5")

        self.page.select_option("#presetSelect", "日曜4人")
        self.assertEqual(self.page.locator("#sundayDayRequired").input_value(), "4")
        self.assertEqual(self.dialogs, ["プリセット名を入力してください:"])


if __name__ == "__main__":
    unittest.main()
